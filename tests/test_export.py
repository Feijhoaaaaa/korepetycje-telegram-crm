from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.services.export import build_account_workbook, build_full_workbook, build_workbook


def test_workbook_contains_localized_lesson() -> None:
    lesson = SimpleNamespace(
        lesson_datetime=datetime(2026, 9, 24, 15, 0, tzinfo=UTC),
        study_chat=SimpleNamespace(student_name="Александр"),
        status="completed",
        payment_status="paid",
    )
    content = build_workbook([lesson], ZoneInfo("Europe/Warsaw"))
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["Занятия"]

    assert [cell.value for cell in sheet[1]] == ["Дата", "Время", "Ученик", "Статус", "Оплата"]
    assert sheet["C2"].value == "Александр"
    assert sheet["D2"].value == "Проведено"
    assert sheet["E2"].value == "Да"
    assert sheet.freeze_panes == "A2"


def test_full_and_student_exports_have_required_sheets() -> None:
    student = SimpleNamespace(name="Анна", teacher=SimpleNamespace(name="Иван"))
    lesson = SimpleNamespace(
        lesson_datetime=datetime(2026, 9, 24, 15, 0, tzinfo=UTC),
        student=student,
        teacher=student.teacher,
        status="completed",
        duration_minutes=60,
        price=100,
        settlement=SimpleNamespace(),
    )
    receipt = SimpleNamespace(
        created_at=datetime(2026, 9, 23, 15, 0, tzinfo=UTC),
        student=student,
        status="approved",
        submitted_amount=150,
        amount=150,
        telegram_file_type="photo",
        rejection_reason=None,
    )
    timezone = ZoneInfo("Europe/Warsaw")
    full = load_workbook(BytesIO(build_full_workbook([lesson], [receipt], [(student, 50)], timezone)))
    account = load_workbook(
        BytesIO(
            build_account_workbook(
                [lesson],
                [receipt],
                [
                    {
                        "datetime": lesson.lesson_datetime,
                        "kind": "lesson",
                        "description": "Lekcja",
                        "amount": -100,
                        "balance": 50,
                    }
                ],
                timezone,
            )
        )
    )
    assert full.sheetnames == ["Lekcje", "Wpłaty", "Salda"]
    assert account.sheetnames == ["Historia konta", "Lekcje", "Wpłaty"]
