from __future__ import annotations

from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.database.models import Lesson, PaymentReceipt, Student
from app.utils.dates import utc_to_local

STATUS = {"planned": "Запланировано", "completed": "Проведено", "cancelled": "Отменено"}
PAYMENT = {"unknown": "—", "paid": "Да", "unpaid": "Нет"}


def build_workbook(lessons: list[Lesson], timezone: ZoneInfo) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Занятия"
    sheet.append(["Дата", "Время", "Ученик", "Статус", "Оплата"])
    for lesson in lessons:
        local = utc_to_local(lesson.lesson_datetime, timezone)
        sheet.append(
            [
                local.date(),
                local.time().replace(tzinfo=None),
                getattr(lesson, "student_name_snapshot", None)
                or (lesson.study_chat.student_name if lesson.study_chat else None)
                or "—",
                STATUS[lesson.status],
                PAYMENT[lesson.payment_status],
            ]
        )
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 13
    sheet.column_dimensions["B"].width = 10
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 18
    sheet.column_dimensions["E"].width = 14
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 1).number_format = "DD.MM.YYYY"
        sheet.cell(row, 2).number_format = "HH:MM"
    for column in range(1, 6):
        sheet.column_dimensions[get_column_letter(column)].bestFit = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(1, sheet.max_row + 1)]
        sheet.column_dimensions[get_column_letter(column)].width = min(max(map(len, values)) + 2, 42)


def _lesson_sheet(workbook: Workbook, lessons: list[Lesson], timezone: ZoneInfo):
    sheet = (
        workbook.active
        if workbook.active.title == "Sheet" and workbook.active.max_row == 1 and workbook.active["A1"].value is None
        else workbook.create_sheet()
    )
    sheet.title = "Lekcje"
    sheet.append(["Data", "Godzina", "Uczeń", "Wykładowca", "Status", "Czas (min)", "Cena PLN", "Opłacona"])
    for lesson in sorted(lessons, key=lambda row: row.lesson_datetime):
        local = utc_to_local(lesson.lesson_datetime, timezone)
        sheet.append(
            [
                local.date(),
                local.time().replace(tzinfo=None),
                lesson.student.name
                if lesson.student
                else getattr(lesson, "student_name_snapshot", None) or "—",
                lesson.teacher.name
                if lesson.teacher
                else getattr(lesson, "teacher_name_snapshot", None) or "—",
                STATUS[lesson.status],
                lesson.duration_minutes,
                lesson.price,
                "Tak" if lesson.settlement or lesson.payment_status == "paid" else "Nie",
            ]
        )
    _style(sheet)
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 1).number_format = "DD.MM.YYYY"
        sheet.cell(row, 2).number_format = "HH:MM"
        sheet.cell(row, 7).number_format = '0.00 "PLN"'
    return sheet


def _payment_sheet(workbook: Workbook, receipts: list[PaymentReceipt]):
    sheet = workbook.create_sheet("Wpłaty")
    sheet.append(
        ["Data", "Uczeń", "Status", "Kwota zgłoszona", "Kwota zatwierdzona", "Waluta", "Typ pliku", "Powód odrzucenia"]
    )
    for row in sorted(receipts, key=lambda item: item.created_at):
        sheet.append(
            [
                row.created_at.replace(tzinfo=None),
                row.student.name,
                row.status,
                row.submitted_amount,
                row.amount,
                getattr(row, "currency", "PLN"),
                row.telegram_file_type,
                row.rejection_reason,
            ]
        )
    _style(sheet)
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 1).number_format = "DD.MM.YYYY HH:MM"
        sheet.cell(row, 4).number_format = sheet.cell(row, 5).number_format = "0.00"
    return sheet


def _save(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_full_workbook(
    lessons: list[Lesson],
    receipts: list[PaymentReceipt],
    balances: list[tuple[Student, object]],
    timezone: ZoneInfo,
) -> bytes:
    workbook = Workbook()
    _lesson_sheet(workbook, lessons, timezone)
    _payment_sheet(workbook, receipts)
    sheet = workbook.create_sheet("Salda")
    sheet.append(["Uczeń", "Wykładowca", "Saldo PLN"])
    for student, balance in sorted(balances, key=lambda item: item[0].name):
        sheet.append([student.name, student.teacher.name, balance])
    _style(sheet)
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 3).number_format = '0.00 "PLN"'
    return _save(workbook)


def build_account_workbook(
    lessons: list[Lesson],
    receipts: list[PaymentReceipt],
    history: list[dict],
    timezone: ZoneInfo,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Historia konta"
    sheet.append(["Data", "Rodzaj", "Opis", "Kwota PLN", "Saldo PLN"])
    for row in reversed(history):
        sheet.append(
            [
                row["datetime"].replace(tzinfo=None),
                row["kind"],
                row["description"],
                row["amount"],
                row["balance"],
            ]
        )
    _style(sheet)
    for index in range(2, sheet.max_row + 1):
        sheet.cell(index, 1).number_format = "DD.MM.YYYY HH:MM"
        sheet.cell(index, 4).number_format = sheet.cell(index, 5).number_format = '0.00 "PLN"'
    _lesson_sheet(workbook, lessons, timezone)
    _payment_sheet(workbook, receipts)
    return _save(workbook)
